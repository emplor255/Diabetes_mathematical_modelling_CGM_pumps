"""Paired, participant-level summaries of the existing control simulations."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .data import STATUS_ORDER


BASELINE = "No external control"
OPTIMIZED = "Optimized unitless control"
# Direction is used only to count improvement; effort is a resource, not an outcome.
METRICS = (
    ("time_70_110_pct", "Time in 70-110 mg/dL", "%", 1),
    ("time_above_110_pct", "Time above 110 mg/dL", "%", -1),
    ("time_below_70_pct", "Time below 70 mg/dL", "%", -1),
    ("peak_glucose", "Peak glucose", "mg/dL", -1),
    ("paper_objective", "Paper objective J", "model objective units", -1),
    ("total_control_intensity", "Integrated control intensity", "model intensity x h", 0),
)
TIE_ATOL = 1e-6


@dataclass
class ControlEvaluation:
    statistics: pd.DataFrame
    paired: pd.DataFrame
    diagnostics: pd.DataFrame
    bootstrap_seed: int
    bootstrap_resamples: int


def evaluate_control_cohort(
    summary: pd.DataFrame,
    participants: pd.DataFrame,
    seed: int = 20260815,
    bootstrap_resamples: int = 10000,
) -> ControlEvaluation:
    """Retain every fitted participant, match by ID, and bootstrap paired changes.

    No optimizer-failure or outcome-based exclusions are made. Missing or duplicate
    scenario rows are errors rather than silently changing the analysis cohort.
    The saved, uniform-grid metrics are reused exactly; this does not rerun control.
    """
    columns = [metric[0] for metric in METRICS]
    required = {"userID", "status", "scenario", "optimizer_success", *columns,
                "min_glucose", "time_at_control_upper_bound_pct"}
    missing = required.difference(summary.columns)
    if missing:
        raise ValueError(f"Control summary is missing columns: {sorted(missing)}")
    if participants.empty or participants["userID"].duplicated().any():
        raise ValueError("Expected a nonempty table of unique fitted participants")
    if bootstrap_resamples < 1:
        raise ValueError("bootstrap_resamples must be positive")
    if summary.duplicated(["userID", "scenario"]).any():
        raise ValueError("Duplicate participant/scenario control rows")
    if set(summary["scenario"]) != {BASELINE, OPTIMIZED}:
        raise ValueError("Expected exactly the no-control and optimized scenarios")

    expected = participants.set_index("userID").sort_index()
    before = summary.loc[summary["scenario"].eq(BASELINE)].set_index("userID")
    after = summary.loc[summary["scenario"].eq(OPTIMIZED)].set_index("userID")
    if set(before.index) != set(expected.index) or set(after.index) != set(expected.index):
        raise ValueError("Both control scenarios must cover every fitted participant exactly")
    before, after = before.loc[expected.index], after.loc[expected.index]
    if not (before["status"].eq(expected["status"]).all()
            and after["status"].eq(expected["status"]).all()):
        raise ValueError("Control and fitted-participant status labels disagree")
    if not set(expected["status"]).issubset(STATUS_ORDER):
        raise ValueError("Unrecognized participant status")
    numeric = columns + ["min_glucose", "time_at_control_upper_bound_pct"]
    for frame in (before, after):
        if not np.isfinite(frame[numeric].to_numpy(dtype=float)).all():
            raise ValueError("Nonfinite control metrics; participants cannot be silently dropped")
        percentages = frame[["time_70_110_pct", "time_above_110_pct", "time_below_70_pct"]]
        if ((percentages < 0) | (percentages > 100)).any().any():
            raise ValueError("Control time percentages must lie within [0, 100]")
        if not np.allclose(percentages.sum(axis=1), 100.0, atol=1e-6):
            raise ValueError("Below/in/above target percentages must sum to 100")

    paired = pd.DataFrame({"userID": expected.index, "status": expected["status"].to_numpy()})
    for metric in columns:
        paired[f"{metric}_no_control"] = before[metric].to_numpy(dtype=float)
        paired[f"{metric}_optimized"] = after[metric].to_numpy(dtype=float)
        paired[f"{metric}_delta"] = (after[metric] - before[metric]).to_numpy(dtype=float)

    groups = [("All participants", np.ones(len(expected), dtype=bool))]
    groups.extend((status, expected["status"].eq(status).to_numpy()) for status in STATUS_ORDER)
    statistic_rows, diagnostic_rows = [], []
    for group_index, (group, mask) in enumerate(groups):
        n = int(mask.sum())
        if n == 0:
            continue
        first, second = before.loc[mask], after.loc[mask]
        # Resample people, preserving each person's before/after pair and all metrics.
        resample_indices = np.random.default_rng(seed + group_index).integers(
            0, n, size=(bootstrap_resamples, n)
        )
        for metric, label, unit, direction in METRICS:
            x, y = first[metric].to_numpy(dtype=float), second[metric].to_numpy(dtype=float)
            delta = y - x
            boot_means = delta[resample_indices].mean(axis=1)
            ci_low, ci_high = np.quantile(boot_means, [0.025, 0.975]) if n > 1 else (np.nan, np.nan)
            row = {
                "group": group, "n_participants": n, "metric": metric,
                "label": label, "unit": unit,
                "delta_unit": "percentage points" if unit == "%" else unit,
                "preferred_direction": {1: "higher", -1: "lower", 0: "resource"}[direction],
                "delta_mean": float(delta.mean()), "delta_ci95_low": float(ci_low),
                "delta_ci95_high": float(ci_high),
                "n_decreased": int((delta < -TIE_ATOL).sum()),
                "n_unchanged": int((np.abs(delta) <= TIE_ATOL).sum()),
                "n_increased": int((delta > TIE_ATOL).sum()),
                "n_improved": int((direction * delta > TIE_ATOL).sum()) if direction else np.nan,
                "n_worsened": int((direction * delta < -TIE_ATOL).sum()) if direction else np.nan,
            }
            for prefix, values in (("no_control", x), ("optimized", y), ("delta", delta)):
                row.update({
                    f"{prefix}_mean": float(values.mean()),
                    f"{prefix}_sd": float(values.std(ddof=1)) if n > 1 else np.nan,
                    f"{prefix}_median": float(np.median(values)),
                    f"{prefix}_q1": float(np.quantile(values, 0.25)),
                    f"{prefix}_q3": float(np.quantile(values, 0.75)),
                })
            statistic_rows.append(row)
        success = second["optimizer_success"].astype(str).str.lower()
        if not success.isin(["true", "false"]).all():
            raise ValueError("optimizer_success must contain booleans")
        low_before, low_after = first["time_below_70_pct"].gt(0), second["time_below_70_pct"].gt(0)
        diagnostic_rows.append({
            "group": group, "n_participants": n,
            "optimizer_successes": int(success.eq("true").sum()),
            "participants_below70_no_control": int(low_before.sum()),
            "participants_below70_optimized": int(low_after.sum()),
            "new_participants_below70": int((~low_before & low_after).sum()),
            "minimum_glucose_no_control": float(first["min_glucose"].min()),
            "minimum_glucose_optimized": float(second["min_glucose"].min()),
            "participants_at_control_upper_bound": int(second["time_at_control_upper_bound_pct"].gt(0).sum()),
        })
    return ControlEvaluation(pd.DataFrame(statistic_rows), paired, pd.DataFrame(diagnostic_rows),
                             seed, bootstrap_resamples)


def presentation_table(statistics: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for row in statistics.itertuples(index=False):
        resource = row.preferred_direction == "resource"
        counts = "--" if resource else f"{int(row.n_improved)}/{row.n_unchanged}/{int(row.n_worsened)}"
        ci = (f"{row.delta_mean:+.2f} [{row.delta_ci95_low:+.2f}, {row.delta_ci95_high:+.2f}]"
              if np.isfinite(row.delta_ci95_low) else f"{row.delta_mean:+.2f} [not estimated]")
        rows.append({
            "Group": row.group, "N": row.n_participants,
            "Outcome (unit)": f"{row.label} ({row.unit})",
            "No control, mean (SD)": f"{row.no_control_mean:.2f} ({row.no_control_sd:.2f})",
            "Optimized, mean (SD)": f"{row.optimized_mean:.2f} ({row.optimized_sd:.2f})",
            "Paired change [95% CI]": ci,
            "Improved / unchanged / worsened": counts,
        })
    return pd.DataFrame(rows)


def _markdown_table(table: pd.DataFrame) -> list[str]:
    # Avoid an optional tabulate dependency.
    return ["| " + " | ".join(table.columns) + " |",
            "| " + " | ".join(["---"] * len(table.columns)) + " |",
            *["| " + " | ".join(str(cell) for cell in row) + " |"
              for row in table.itertuples(index=False, name=None)]]


def control_table_notes(evaluation: ControlEvaluation, audit: dict[str, Any]) -> list[str]:
    n = len(evaluation.paired)
    fit_events = audit.get("fit_events", "the usable")
    horizon = audit.get("control_config", {}).get("horizon_hours", 2.5)
    step = audit.get("control_config", {}).get("integration_step_minutes", 1.0)
    return [
        f"Coverage: all {n} fitted participants, with one {horizon:g}-hour simulated meal scenario per person. "
        "Each scenario uses that person's median fitted meal amplitude and delay, starting at reference glucose. "
        f"The {fit_events} fitted meal windows are used for parameter estimation; this table does not evaluate "
        "each observed meal separately or include Hall2018 participants without fitted meal data.",
        "Scenario values are participant means (sample SD), with equal weight per person. "
        "Change is optimized minus no control, computed within each person before aggregation. "
        "Time-percentage changes are percentage points, not relative percentages. "
        f"Time percentages reuse the saved fractions of simulated samples on the {step:g}-minute grid, "
        "including both endpoints, to match the existing control results.",
        f"95% CIs are percentile bootstrap intervals for the mean paired change, resampling participant pairs "
        f"{evaluation.bootstrap_resamples:,} times (seed {evaluation.bootstrap_seed}). "
        "They describe between-participant sampling uncertainty conditional on fitted parameters; "
        "they do not propagate parameter, model, or optimizer uncertainty. Small subgroups are descriptive. "
        "No p-values are used to claim superiority.",
        "Improved/unchanged/worsened counts follow the favorable direction for each outcome: higher time "
        "in target; lower time above/below target, peak glucose, and objective. The numerical tie tolerance "
        f"is {TIE_ATOL:g} in the metric's units. Peak reduction must be interpreted together with time below target. "
        "Control intensity is a resource metric with no improvement classification.",
        "The paper objective may be negative; its signed paired change in original units is reported. "
        "It equals a time-weighted squared distance from the target midpoint plus effort, up to a constant. "
        "Reducing it need not improve time in target or prevent glucose below 70 mg/dL. "
        "Because this is the quantity optimized, its reduction alone is not independent validation.",
        "The comparator is no external control. These are in-model, open-loop simulations with the same "
        "fitted dynamics used for optimization, not held-out predictions or observed treatment outcomes. "
        "The table does not establish superiority to another controller or clinical efficacy. "
        "Control intensity has model units, with no calibration to administered insulin dose.",
    ]


def control_report_lines(evaluation: ControlEvaluation, audit: dict[str, Any]) -> list[str]:
    stats = evaluation.statistics
    overall = presentation_table(stats.loc[stats["group"].eq("All participants")]).drop(columns=["Group", "N"])
    lines = ["### Cohort-wide paired control comparison", "",
             f"Table 1. All {len(evaluation.paired)} fitted participants; no outcome-based exclusions.", "",
             *_markdown_table(overall), ""]
    diag = evaluation.diagnostics.iloc[0]
    lines.extend([
        f"Participants with any simulated glucose below 70 mg/dL: "
        f"{int(diag.participants_below70_no_control)}/{int(diag.n_participants)} without control versus "
        f"{int(diag.participants_below70_optimized)}/{int(diag.n_participants)} with control "
        f"({int(diag.new_participants_below70)} newly affected). "
        f"Minimum simulated glucose across participants: {diag.minimum_glucose_no_control:.2f} versus "
        f"{diag.minimum_glucose_optimized:.2f} mg/dL. "
        f"Optimizers reported success for {int(diag.optimizer_successes)}/{int(diag.n_participants)}; "
        f"{int(diag.participants_at_control_upper_bound)}/{int(diag.n_participants)} reached the upper control bound.",
        "", "### Breakdown by participant status", "",
    ])
    for group in STATUS_ORDER:
        subset = stats.loc[stats["group"].eq(group)]
        if subset.empty:
            continue
        lines.extend([f"{group} (n={int(subset.iloc[0].n_participants)}).", "",
                      *_markdown_table(presentation_table(subset).drop(columns=["Group", "N"])), ""])
    lines.extend(["Table definitions and scope:", "", *[f"- {note}" for note in control_table_notes(evaluation, audit)],
                  "", "Exports: [Excel tables](tables/control_evaluation.xlsx), "
                  "[numeric cohort statistics](tables/control_cohort_statistics.csv), "
                  "[paired participant changes](tables/control_paired_changes.csv), "
                  "[LaTeX table](tables/control_evaluation.tex)."])
    return lines


def _latex_escape(value: str) -> str:
    mapping = {"&": r"\&", "%": r"\%", "_": r"\_", "#": r"\#"}
    return "".join(mapping.get(character, character) for character in value)


def write_control_tables(output_dir: Path, evaluation: ControlEvaluation, audit: dict[str, Any]) -> None:
    directory = output_dir / "tables"
    directory.mkdir(parents=True, exist_ok=True)
    stats = evaluation.statistics
    stats.to_csv(directory / "control_cohort_statistics.csv", index=False)
    evaluation.paired.to_csv(directory / "control_paired_changes.csv", index=False)
    evaluation.diagnostics.to_csv(directory / "control_cohort_diagnostics.csv", index=False)
    display = presentation_table(stats)
    notes = control_table_notes(evaluation, audit)
    with pd.ExcelWriter(directory / "control_evaluation.xlsx", engine="openpyxl") as writer:
        for name, frame in (
            ("Overall", display.loc[display["Group"].eq("All participants")]),
            ("By status", display.loc[~display["Group"].eq("All participants")]),
            ("Numeric statistics", stats), ("Paired changes", evaluation.paired),
            ("Diagnostics", evaluation.diagnostics), ("Read me", pd.DataFrame({"Notes": notes})),
        ):
            frame.to_excel(writer, sheet_name=name, index=False)
            worksheet = writer.sheets[name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="24475B")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for column_index, column in enumerate(frame.columns, start=1):
                width = min(54, max(14, len(str(column)) + 2,
                                   *(len(str(value)) + 2 for value in frame[column].head(40))))
                worksheet.column_dimensions[get_column_letter(column_index)].width = width
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            if name == "Read me":
                worksheet.column_dimensions["A"].width = 120
                for row_index in range(2, len(notes) + 2):
                    worksheet.row_dimensions[row_index].height = 80

    # A plain LaTeX tabular: no booktabs, pandas styling, or extra TeX package needed.
    n = len(evaluation.paired)
    tex = [r"\begin{table}[htbp]", r"\centering", r"\footnotesize",
           f"\\caption{{Paired comparison of simulated control outcomes across all {n} fitted participants.}}",
           r"\label{tab:control-cohort}",
           r"\begin{tabular}{p{.22\linewidth}p{.14\linewidth}p{.14\linewidth}p{.24\linewidth}p{.13\linewidth}}",
           r"\hline",
           r"Outcome & No control & Optimized & Change [95\% CI] & Better/same/worse \\", r"\hline"]
    for row in display.loc[display["Group"].eq("All participants")].itertuples(index=False, name=None):
        # Group and N are already in the caption.
        tex.append(" & ".join(_latex_escape(str(value)) for value in row[2:]) + r" \\")
    tex.extend([r"\hline", r"\end{tabular}", r"\par\smallskip", r"\begin{minipage}{\textwidth}",
                r"\scriptsize",
                _latex_escape("Values are mean (SD). Change = optimized minus no control; time changes are percentage points. "
                              f"Paired percentile bootstrap CIs use {evaluation.bootstrap_resamples:,} resamples. "
                              "One median-meal simulation per fitted participant; conditional on fitted model parameters. "
                              "Control is uncalibrated model intensity. Full definitions and scope are in the analysis report."),
                r"\end{minipage}", r"\end{table}"])
    (directory / "control_evaluation.tex").write_text("\n".join(tex) + "\n", encoding="utf-8")
